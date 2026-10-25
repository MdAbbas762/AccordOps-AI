import hashlib
import re
import zipfile
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader


MAX_CONTRACT_SIZE_BYTES = 10 * 1024 * 1024
PDF_MIME_TYPE = "application/pdf"
DOCX_MIME_TYPE = (
    "application/vnd.openxmlformats-officedocument."
    "wordprocessingml.document"
)
SUPPORTED_EXTENSIONS = {
    ".pdf": PDF_MIME_TYPE,
    ".docx": DOCX_MIME_TYPE,
}


class DocumentValidationError(ValueError):
    pass


@dataclass(frozen=True)
class ParsedDocument:
    text: str
    page_count: int | None
    mime_type: str
    file_hash: str
    safe_filename: str


def sanitize_filename(filename: str) -> str:
    raw_name = Path(filename).name
    suffix = Path(raw_name).suffix.lower()
    stem = Path(raw_name).stem
    clean_stem = re.sub(r"[^A-Za-z0-9._-]+", "_", stem).strip("._")
    if not clean_stem:
        clean_stem = "contract"
    return f"{clean_stem[:120]}{suffix}"


def _extract_pdf(content: bytes) -> tuple[str, int]:
    if not content.startswith(b"%PDF"):
        raise DocumentValidationError(
            "The selected file does not appear to be a valid PDF."
        )

    try:
        reader = PdfReader(BytesIO(content))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:
                raise DocumentValidationError(
                    "Password-protected PDF files are not supported."
                ) from exc

        text_parts: list[str] = []
        for index, page in enumerate(reader.pages, start=1):
            page_text = (page.extract_text() or "").strip()
            if page_text:
                text_parts.append(f"[Page {index}]\n{page_text}")

        return "\n\n".join(text_parts).strip(), len(reader.pages)
    except DocumentValidationError:
        raise
    except Exception as exc:
        raise DocumentValidationError("The PDF could not be read.") from exc


def _extract_docx(content: bytes) -> tuple[str, None]:
    if not zipfile.is_zipfile(BytesIO(content)):
        raise DocumentValidationError(
            "The selected file does not appear to be a valid DOCX document."
        )

    try:
        document = Document(BytesIO(content))
    except Exception as exc:
        raise DocumentValidationError(
            "The DOCX document could not be read."
        ) from exc

    text_parts: list[str] = []
    for paragraph in document.paragraphs:
        value = paragraph.text.strip()
        if value:
            text_parts.append(value)

    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if values:
                text_parts.append(" | ".join(values))

    return "\n".join(text_parts).strip(), None


def parse_contract_document(filename: str, content: bytes) -> ParsedDocument:
    if not filename:
        raise DocumentValidationError("A file name is required.")
    if not content:
        raise DocumentValidationError("The uploaded file is empty.")
    if len(content) > MAX_CONTRACT_SIZE_BYTES:
        raise DocumentValidationError("Contract files must be 10 MB or smaller.")

    extension = Path(filename).suffix.lower()
    mime_type = SUPPORTED_EXTENSIONS.get(extension)
    if mime_type is None:
        raise DocumentValidationError(
            "Only PDF and DOCX contract files are supported."
        )

    if extension == ".pdf":
        text, page_count = _extract_pdf(content)
    else:
        text, page_count = _extract_docx(content)

    if not text.strip():
        raise DocumentValidationError(
            "No extractable text was found. Scanned/image-only documents "
            "are not supported in this MVP."
        )

    return ParsedDocument(
        text=text,
        page_count=page_count,
        mime_type=mime_type,
        file_hash=hashlib.sha256(content).hexdigest(),
        safe_filename=sanitize_filename(filename),
    )
