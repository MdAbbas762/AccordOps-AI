from io import BytesIO

import pytest
from docx import Document

from app.services.documents import (
    DocumentValidationError,
    parse_contract_document,
    sanitize_filename,
)


def test_sanitize_filename() -> None:
    assert (
        sanitize_filename("Customer Contract (Final).PDF")
        == "Customer_Contract_Final.pdf"
    )


def test_parse_docx_contract() -> None:
    document = Document()
    document.add_heading("Service Agreement", level=1)
    document.add_paragraph(
        "Provider must respond to priority-one incidents within 30 minutes."
    )

    output = BytesIO()
    document.save(output)

    result = parse_contract_document(
        "agreement.docx",
        output.getvalue(),
    )

    assert "Service Agreement" in result.text
    assert "30 minutes" in result.text
    assert result.page_count is None
    assert len(result.file_hash) == 64


def test_reject_unsupported_document() -> None:
    with pytest.raises(DocumentValidationError):
        parse_contract_document(
            "contract.txt",
            b"not supported",
        )
